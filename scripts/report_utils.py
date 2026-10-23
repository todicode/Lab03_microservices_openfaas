from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import math
_PAGE_BREAK=False

ROOT=Path(__file__).resolve().parents[1]
REPORTS=ROOT/"reports"
REPORTS.mkdir(exist_ok=True)
SHOTS=ROOT/"evidence/report-shots"

def document():
    d=Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.75)
    sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
    for name in ("Normal","Title","Subtitle","Heading 1","Heading 2","Caption"):
        style=d.styles[name]
        style.font.name="Times New Roman"; style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"Times New Roman")
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower(): del fonts.attrib[attr]
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    normal=d.styles["Normal"]
    normal.font.size=Pt(12); normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(6)
    for name,size in (("Heading 1",16),("Heading 2",13)):
        s=d.styles[name]; s.font.size=Pt(size); s.font.bold=True
        s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(7)
        s.paragraph_format.keep_with_next=True
    d.styles["Title"].font.size=Pt(25)
    d.styles["Title"].font.bold=True
    d.styles["Subtitle"].font.size=Pt(16)
    d.styles["Caption"].font.size=Pt(10.5)
    d.styles["Caption"].paragraph_format.space_after=Pt(10)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run(); r.font.name="Times New Roman"; r.font.size=Pt(10)
    f=OxmlElement("w:fldSimple"); f.set(qn("w:instr"),"PAGE")
    r._r.addnext(f)
    d.core_properties.title="Báo cáo Lab 03 Triển khai microservices bằng OpenFaaS"
    d.core_properties.subject="Ứng dụng Paperlane với Python PostgreSQL và Kubernetes"
    d.core_properties.author=""
    return d

def para(d,text,bold=None):
    p=d.add_paragraph(); r=p.add_run(text); r.bold=bold
    return p

def heading(d,text,level=1):
    global _PAGE_BREAK
    p=d.add_heading(text,level)
    if _PAGE_BREAK:
        p.paragraph_format.page_break_before=True
        _PAGE_BREAK=False
    return p

def new_page(d):
    global _PAGE_BREAK
    _PAGE_BREAK=True

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,text,w in zip(t.rows[0].cells,headers,widths):
        c.width=Inches(w); c.text=text
    repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,value,w in zip(cells,row,widths):
            c.width=Inches(w); c.text=str(value)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); no_split=OxmlElement("w:cantSplit"); trpr.append(no_split)
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"E8EDF2" if i==0 else ("F7F8FA" if i%2==0 else "FFFFFF")); tcpr.append(shade)
            margins=OxmlElement("w:tcMar")
            for edge in ("top","bottom","left","right"):
                element=OxmlElement("w:"+edge); element.set(qn("w:w"),"60"); element.set(qn("w:type"),"dxa"); margins.append(element)
            tcpr.append(margins)
            borders=OxmlElement("w:tcBorders")
            for edge in ("top","left","bottom","right"):
                e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
            tcpr.append(borders)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
                for r in p.runs:
                    r.font.name="Times New Roman";r.font.size=Pt(10.5);r.font.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def code(d,text):
    for line in text.splitlines():
        p=d.add_paragraph()
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        r=p.add_run(line);r.font.name="Consolas";r.font.size=Pt(9.5)
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def picture(d,name,caption,width=6.65):
    path=SHOTS/name
    if not path.exists(): raise FileNotFoundError(path)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    inline=p.runs[0]._r.xpath(".//wp:docPr")[0]
    inline.set("descr",caption)
    cap=d.add_paragraph(caption,"Caption");cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

def architecture():
    target=REPORTS/"architecture.png"
    im=Image.new("RGB",(1900,1080),"white"); dr=ImageDraw.Draw(im)
    font=lambda s: ImageFont.truetype("C:/Windows/Fonts/arial.ttf",s)
    def box(rect,title,sub="",fill="#eef3ef"):
        dr.rounded_rectangle(rect,radius=17,fill=fill,outline="#536b59",width=3)
        x1,y1,x2,y2=rect
        dr.text(((x1+x2)/2,(y1+y2)/2-15 if sub else (y1+y2)/2),title,font=font(30),fill="#17281d",anchor="mm")
        if sub:dr.text(((x1+x2)/2,(y1+y2)/2+29),sub,font=font(23),fill="#425b49",anchor="mm")
    def arrow(points,color="#344b3a",width=4):
        dr.line(points,fill=color,width=width)
        a,b=points[-2],points[-1];ang=math.atan2(b[1]-a[1],b[0]-a[0])
        tri=[b,(b[0]-18*math.cos(ang-.45),b[1]-18*math.sin(ang-.45)),(b[0]-18*math.cos(ang+.45),b[1]-18*math.sin(ang+.45))]
        dr.polygon(tri,fill=color)
    dr.rounded_rectangle((360,60,1840,1000),radius=26,fill="#fafbf9",outline="#b5c0b4",width=3)
    dr.text((400,105),"Kubernetes  lab03-openfaas",font=font(34),fill="black")
    box((35,420,300,570),"Trình duyệt","Frontend",fill="#f4eee2")
    box((430,400,765,585),"OpenFaaS Gateway","Định tuyến HTTP")
    box((970,210,1310,340),"Product Service","Catalog và giá bán")
    box((970,430,1310,560),"Order Service","Điều phối đặt hàng")
    box((970,650,1310,780),"Total Calculator","Tính tổng tiền VND")
    box((1460,195,1775,345),"Product DB","role  product_app",fill="#eef1f8")
    box((1460,425,1775,575),"Order DB","role  order_app",fill="#eef1f8")
    dr.rounded_rectangle((1420,150,1810,800),radius=15,outline="#8d9fb7",width=2)
    dr.text((1615,650),"PostgreSQL 17",font=font(29),fill="#293f60",anchor="mm")
    dr.text((1615,700),"PVC  1 GiB",font=font(25),fill="#293f60",anchor="mm")
    arrow([(300,495),(430,495)])
    arrow([(765,470),(870,470),(870,275),(970,275)])
    arrow([(765,495),(970,495)])
    arrow([(765,520),(870,520),(870,715),(970,715)])
    arrow([(1310,275),(1460,275)])
    arrow([(1310,495),(1460,495)])
    arrow([(970,540),(920,540),(920,880),(595,880),(595,585)],color="#92794b",width=4)
    dr.text((940,855),"Order gọi Product và Calculator qua Gateway",font=font(26),fill="#735c32",anchor="mm")
    dr.text((400,955),"Secret kết nối được mount vào từng function  •  Database nằm ngoài container function",font=font(25),fill="#465046")
    im.save(target)
    return target
